# 3단계 ② — fetch_post_files.js 가 받아 둔 파일(캐시)과 글 본문 표를 읽어 **건별 줄**로 푼다.
#
#   python scripts/read_post_files.py            data/post_files.json 전부
#   python scripts/read_post_files.py --파일 X   파일 하나만 읽어 표를 찍는다(시험)
#
# 읽는 꼴 — xlsx·xls·csv·pdf·hwpx·hwp·zip(안의 것), 그리고 글 본문의 HTML 표.
#   hwp 는 한글 5.0 바이너리를 직접 푼다(olefile). **배포용 문서(복사 막음)·암호 걸린 것은 못 읽는다** — 그렇다고 적는다.
# 푸는 법 — 표마다 머리줄(일시·장소·목적·인원·금액·방법 가운데 둘 이상)을 찾아 칸을 맞추고,
#   그 아래에서 **금액이 있는 줄**만 건으로 친다. 합계·소계 줄은 뺀다. 날짜에 해가 없으면 글 제목(「2026년 9월」)·글 날짜에서 채운다.
#   ponytail: 머리줄 낱말 맞추기라 칸 이름이 아주 다른 곳은 「표 못 찾음」으로 남는다 — 그런 곳이 많으면 낱말을 늘린다.
#
# 결과
#   data/expense_rows.json.gz — 업무추진비(단체장·의장) 건별 줄 {곳, 이름, 고르기, 글, 날짜, 장소, 내용, 인원, 금액, 방법}
#   data/expense_summary.json — 곳마다 기간·건수·합계·카드/현금·경조사·많이 쓴 곳 다섯(회의 재료, 판정 아님)
#   data/post_read.json       — 글마다 읽은 결과(줄 수·합계·파일별 형식과 까닭). 못 읽은 것을 세는 데 쓴다
#   data/trip_reports.json    — 의회 국외출장: 해마다 계획 글·결과 글 수(제목으로 가름). 글마다 받은 파일 수·pdf 쪽수는 post_read.json

import datetime
import gzip
import warnings
import io
import json
import os
import re
import struct
import sys
import zipfile
import zlib

try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except AttributeError:
    pass

warnings.filterwarnings('ignore')   # openpyxl 머리글·인쇄 영역 경고 — 값 읽기와 무관
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
캐시 = os.environ.get('LF_FILES') or os.path.join(os.environ.get('LOCALAPPDATA') or os.path.expanduser('~'), 'local-finance-files')


# ── 꼴마다 표 꺼내기: 모두 [[칸, 칸, …], …] 의 목록을 돌려준다 ──

def 칸글(v):
    if v is None:
        return ''
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d %H:%M') if (v.hour or v.minute) else v.strftime('%Y-%m-%d')
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r'\s+', ' ', str(v)).strip()


def 표_xlsx(몸):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(몸), read_only=True, data_only=True)
    return [[[칸글(c) for c in r] for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]


def 표_xls(몸):
    import xlrd
    wb = xlrd.open_workbook(file_contents=몸)
    표들 = []
    for sh in wb.sheets():
        줄들 = []
        for i in range(sh.nrows):
            줄 = []
            for j, c in enumerate(sh.row(i)):
                if c.ctype == xlrd.XL_CELL_DATE:
                    줄.append(칸글(xlrd.xldate_as_datetime(c.value, wb.datemode)))
                else:
                    줄.append(칸글(c.value))
            줄들.append(줄)
        표들.append(줄들)
    return 표들


def 표_csv(몸):
    import csv
    for enc in ('utf-8-sig', 'cp949'):
        try:
            글 = 몸.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        return []
    return [[[칸글(c) for c in r] for r in csv.reader(io.StringIO(글))]]


def 표_pdf(몸):
    import pymupdf
    doc = pymupdf.open(stream=몸, filetype='pdf')
    표들 = []
    for page in doc:
        try:
            for t in page.find_tables().tables:
                표들.append([[칸글(c) for c in r] for r in t.extract()])
        except Exception:
            pass
    if sum(len(p.get_text().strip()) for p in doc) < 30:
        raise 못읽음('그림 pdf(글자 없음)')
    if not 표들:   # 표 선이 없는 pdf — 줄 글자를 띄어쓰기 두 칸 이상으로 갈라 표처럼 본다
        줄들 = [re.split(r'\s{2,}', l.strip()) for p in doc for l in p.get_text().splitlines() if l.strip()]
        표들.append(줄들)
    return 표들


def 표_hwpx(몸):
    z = zipfile.ZipFile(io.BytesIO(몸))
    표들 = []
    for n in sorted(x for x in z.namelist() if re.match(r'Contents/section\d+\.xml', x)):
        x = z.read(n).decode('utf-8', 'replace')
        # 가장 안쪽 표부터 — 표 안의 표는 따로 센다
        for tbl in re.findall(r'<hp:tbl\b(?:(?!<hp:tbl\b).)*?</hp:tbl>', x, re.S):
            줄들 = []
            for tr in re.findall(r'<hp:tr\b.*?</hp:tr>', tbl, re.S):
                줄들.append([칸글(' '.join(re.findall(r'<hp:t>([^<]*)</hp:t>', tc))) for tc in re.findall(r'<hp:tc\b.*?</hp:tc>', tr, re.S)])
            표들.append(줄들)
    return 표들


class 못읽음(Exception):
    pass


def 표_hwp(몸):
    """한글 5.0 바이너리 — BodyText/Section* 의 레코드를 따라가며 표(셀 주소)를 다시 세운다."""
    import olefile
    ole = olefile.OleFileIO(io.BytesIO(몸))
    머리 = ole.openstream('FileHeader').read()
    깃 = struct.unpack('<I', 머리[36:40])[0]
    if 깃 & 2:
        raise 못읽음('암호 걸린 hwp')
    if 깃 & 4:
        raise 못읽음('배포용 hwp(복사 막음)')
    눌림 = 깃 & 1
    표들 = []
    섹션 = sorted((s for s in ole.listdir() if s[0] == 'BodyText'), key=lambda s: int(re.sub(r'\D', '', s[1]) or 0))
    for s in 섹션:
        b = ole.openstream(s).read()
        if 눌림:
            b = zlib.decompress(b, -15)
        표들 += _hwp_표(b)
    return 표들


def _hwp_글(b):
    out, i = [], 0
    while i + 1 < len(b):
        c = struct.unpack('<H', b[i:i + 2])[0]
        if c < 32:
            if c in (0, 10, 13) or 24 <= c <= 31:
                out.append(' ' if c in (10, 13) else '')
                i += 2
            else:   # 확장·인라인 조종 문자는 8글자(16바이트) 자리를 차지한다. 9(탭)도 그렇다
                out.append(' ' if c == 9 else '')
                i += 16
        else:
            out.append(chr(c))
            i += 2
    return ''.join(out)


def _hwp_표(b):
    레코드 = []
    i = 0
    while i + 4 <= len(b):
        h = struct.unpack('<I', b[i:i + 4])[0]
        tag, lv, sz = h & 0x3FF, (h >> 10) & 0x3FF, (h >> 20) & 0xFFF
        i += 4
        if sz == 0xFFF:
            sz = struct.unpack('<I', b[i:i + 4])[0]
            i += 4
        레코드.append((tag, lv, b[i:i + sz]))
        i += sz
    CTRL, LIST, TABLE, TEXT = 0x47, 0x48, 0x4D, 0x43
    표들, 쌓임 = [], []   # 쌓임: [{lv: 표 조종 레코드 수준, 칸: {(행,열): 글}, 지금: (행,열)}]
    for tag, lv, d in 레코드:
        while 쌓임 and lv <= 쌓임[-1]['lv'] and not (tag in (LIST, TABLE) and lv == 쌓임[-1]['lv'] + 1):
            표들.append(쌓임.pop())
        if tag == CTRL and d[:4] == b' lbt':
            쌓임.append({'lv': lv, '칸': {}, '지금': None})
        elif tag == LIST and 쌓임 and lv == 쌓임[-1]['lv'] + 1 and len(d) >= 12:
            열, 행 = struct.unpack('<HH', d[8:12])
            쌓임[-1]['지금'] = (행, 열)
            쌓임[-1]['칸'].setdefault((행, 열), '')
        elif tag == TEXT and 쌓임 and 쌓임[-1]['지금'] is not None:
            k = 쌓임[-1]['지금']
            쌓임[-1]['칸'][k] = (쌓임[-1]['칸'][k] + ' ' + _hwp_글(d)).strip()
    표들 += 쌓임[::-1]
    out = []
    for t in 표들:
        if not t['칸']:
            continue
        행수 = max(r for r, _ in t['칸']) + 1
        열수 = max(c for _, c in t['칸']) + 1
        out.append([[칸글(t['칸'].get((r, c), '')) for c in range(열수)] for r in range(행수)])
    return out


def 표꺼내기(몸, 확장자):
    """(표 목록, 형식) — 확장자를 믿지 않고 앞머리로 가른다(이름이 .bin·.do 로 오는 곳이 있다)."""
    if 몸[:4] == b'PK\x03\x04':
        z = zipfile.ZipFile(io.BytesIO(몸))
        이름들 = z.namelist()
        if any(n.startswith('xl/') for n in 이름들):
            return 표_xlsx(몸), 'xlsx'
        if any(n.startswith('Contents/') for n in 이름들):
            return 표_hwpx(몸), 'hwpx'
        표들 = []
        for n in 이름들:   # 그냥 zip — 안의 것을 차례로
            if not n.endswith('/'):
                try:
                    표들 += 표꺼내기(z.read(n), os.path.splitext(n)[1][1:])[0]
                except Exception:
                    pass
        return 표들, 'zip'
    if 몸[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
        import olefile
        ole = olefile.OleFileIO(io.BytesIO(몸))
        if ole.exists('FileHeader'):
            return 표_hwp(몸), 'hwp'
        return 표_xls(몸), 'xls'
    if 몸[:5] == b'%PDF-':
        return 표_pdf(몸), 'pdf'
    if 확장자 == 'brf':
        raise 못읽음('점자 파일(건너뜀)')
    if 확장자 == 'csv':
        return 표_csv(몸), 'csv'
    if 몸.lstrip()[:1] == b'<':
        raise 못읽음('파일이 아니라 웹 화면이 받아짐')
    raise 못읽음(f'모르는 꼴({확장자})')


# ── 표를 건별 줄로 ──

칸낱말 = {
    '날짜': r'일시|일자|날짜|사용일|집행일|연월일|년월일|결제일|승인일|사용\s*일시|집행\s*일시',
    '장소': r'장소|사용처|가맹점|업소|상호|집행처|사용장소|업체|지급처',
    '내용': r'목적|내용|사유|용도|건명|내역|적요|집행목적|사용목적',
    '인원': r'대상|인원|참석|인원수|대상자',
    '금액': r'금액|집행액|사용액|결제액|집행금액|사용금액|지출액|지급액|\(원\)|원\)',
    '방법': r'방법|결제|카드|현금',
}
합계줄 = re.compile(r'^\s*(합\s*계|소\s*계|총\s*계|계|누\s*계)\s*$')


def 머리찾기(표):
    """(머리줄 번호, {칸이름: 열번호}) — 한 줄 머리를 먼저 다 보고, 없으면 두 줄 머리(위아래 합친 것).
    현금·카드가 따로 칸인 표(금액 칸이 없다)는 둘을 '현금'·'카드' 로 적어 두고 값이 든 쪽을 금액으로 친다."""
    for 겹 in (1, 2):
        for i, 줄 in enumerate(표[:15]):
            머 = 줄 if 겹 == 1 else [(a + ' ' + b).strip() for a, b in zip(줄, (표[i + 1] if i + 1 < len(표) else []) + [''] * len(줄))]
            맞음 = {}
            for j, c in enumerate(머):
                c = c.replace(' ', '')
                if not c or len(c) > 16:
                    continue
                if re.fullmatch(r'(집행방법|결제방법|사용방법)?(현금|카드)', c):
                    맞음.setdefault(c[-2:], j)
                    continue
                for k, pat in 칸낱말.items():
                    if k not in 맞음 and re.search(pat, c):
                        맞음[k] = j
                        break
            if '금액' not in 맞음 and '현금' in 맞음 and '카드' in 맞음:
                맞음['금액'] = None
            if '금액' in 맞음 and len(set(맞음) - {'금액', '현금', '카드'}) >= 2:
                return i + 겹 - 1, 맞음
    return None, None


def 금액(s):
    s = re.sub(r'[,원\s₩￦]', '', s or '')
    m = re.fullmatch(r'-?\d{3,12}', s)
    return int(m.group()) if m else None


def 날짜(s, 해):
    s = (s or '').strip()
    if re.fullmatch(r'4[4-7]\d{3}(\.\d+)?', s):   # 엑셀 날짜 번호(46113 = 2026-04-01) — 날짜 꼴이 아닌 칸이 숫자로 읽힘(연수구·경기도의회)
        return (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(float(s)))).isoformat()
    m = re.search(r'(20\d{2})\s*[-.\/년]\s*(\d{1,2})\s*[-.\/월]\s*(\d{1,2})', s) or re.search(r'(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)', s)   # 20260318(수성구·고령)
    m2 = None if m else re.search(r'(?<!\d)(\d{2})\s*[-.\/]\s*(\d{1,2})\s*[-.\/]\s*(\d{1,2})(?!\d)', s)   # 25.03.13 — 두 자리 해(단양·세종)
    if m:
        y, mo, d = map(int, m.groups())
    elif m2:
        y, mo, d = map(int, m2.groups())
        y += 2000
    else:
        m = re.search(r'(?<!\d)(\d{1,2})\s*[-.\/월]\s*(\d{1,2})(?!\d)', s)
        if not m or not 해:
            return ''
        y, (mo, d) = 해, map(int, m.groups())
    if not (1 <= mo <= 12 and 1 <= d <= 31):
        return ''
    if 해 and abs(y - 해) > 1:   # 원문 오타 — 「2025년 6월」 글에 2028-06-26, 2055-12-24(안양·파주). 글의 해로
        y = 해
    return f'{y:04d}-{mo:02d}-{d:02d}'


def 표풀기(표, 해):
    i, 칸 = 머리찾기(표)
    if i is None:
        return None
    줄들, 앞날 = [], ''
    for 줄 in 표[i + 1:]:
        if not any(줄):
            continue
        g = lambda k: 줄[칸[k]] if 칸.get(k) is not None and 칸[k] < len(줄) else ''
        날, j = g('날짜').strip(), 칸.get('날짜')
        if re.fullmatch(r'\d{1,2}\.?', 날) and j + 1 < len(줄) and re.fullmatch(r'\d{1,2}\.?', 줄[j + 1].strip()):   # 월·일이 갈린 칸(동해 「사용일시」 아래 월|일|시간)
            날 = f'{날.rstrip(".")}월 {줄[j + 1].strip().rstrip(".")}'
        돈, 방법 = 금액(g('금액')), g('방법')
        if 칸['금액'] is None:   # 현금·카드 갈린 표
            돈 = 금액(g('카드')) or 금액(g('현금'))
            방법 = '카드' if 금액(g('카드')) else '현금'

        if 돈 is None or 돈 <= 0 or any(합계줄.match(c or '') for c in 줄[:3]):
            continue
        if not (날 or g('장소')) or re.fullmatch(r'\d+\s*건', g('내용')):   # 요약표(「50건」)·소계 줄
            continue
        # 날짜 칸이 비면 위 줄 날짜 — 같은 날 건을 한 칸으로 합친 표(세종·단양)
        앞날 = 날짜(날, 해) if 날 else 앞날
        줄들.append({'날짜': 앞날, '장소': g('장소')[:60], '내용': g('내용')[:80],
                     '인원': g('인원')[:30], '금액': 돈, '방법': 방법[:12]})
    return 줄들


def 해찾기(*글자들):
    """날짜에 해가 없을 때 쓸 해 — 앞에 준 것부터(파일 이름 → 글 제목 → 글 올린 날). ‘26. 9월 꼴도 본다."""
    for t in 글자들:
        m = re.search(r'(20\d{2})\s*[년.\-]', t or '') or re.search(r"[‘'’](\d{2})\s*[.년]", t or '')
        if m:
            y = int(m.group(1))
            return y if y > 2000 else 2000 + y
    return None


def 글읽기(글):
    해 = 해찾기(글['제목'], 글.get('날짜'))
    줄들, 파일들 = [], []
    for 표 in 글.get('표') or []:
        r = 표풀기(표, 해)
        if r:
            줄들 += r
    for f in 글.get('파일') or []:
        if not f.get('sha1'):
            continue
        자리 = os.path.join(캐시, f['sha1'][:2], f['sha1'] + '.' + f['확장자'])
        한 = {'sha1': f['sha1'], '이름': f['이름'][:80]}
        try:
            표들, 꼴 = 표꺼내기(open(자리, 'rb').read(), f['확장자'])
            한['꼴'] = 꼴
            파일해 = 해찾기(f['이름'], 글['제목'], 글.get('날짜'))
            풀린 = [r for r in (표풀기(t, 파일해) for t in 표들) if r]
            n = sum(len(r) for r in 풀린)
            한['까닭'] = f'{n}줄' if 풀린 else '표 못 찾음'
            for r in 풀린:
                줄들 += r
        except FileNotFoundError:
            한['까닭'] = '캐시에 없음'
        except 못읽음 as e:
            한['까닭'] = str(e)
        except Exception as e:
            한['까닭'] = '못 읽음 ' + type(e).__name__
        파일들.append(한)
    # 같은 건이 본문 표와 파일에 겹쳐 실리는 곳이 있다 — 날짜·금액·장소가 같으면 하나로
    본, 고유 = set(), []
    for r in 줄들:
        k = (r['날짜'], r['금액'], r['장소'], r['내용'])
        if k not in 본:
            본.add(k)
            고유.append(r)
    return 고유, 파일들


def 출장갈래(제목):
    if re.search(r'결과|보고서|귀국', 제목):
        return '결과'
    if re.search(r'계획|심사|사전|일정', 제목):
        return '계획'
    return '기타'


def 쪽수(f):
    """pdf 만 센다(보고서가 몇 쪽인가). 다른 꼴은 0."""
    try:
        자리 = os.path.join(캐시, f['sha1'][:2], f['sha1'] + '.' + f['확장자'])
        몸 = open(자리, 'rb').read()
        if 몸[:5] == b'%PDF-':
            import pymupdf
            return len(pymupdf.open(stream=몸, filetype='pdf'))
    except Exception:
        pass
    return 0


def 출장셈():
    """의회 국외출장 — 글 **제목**만으로 해마다 계획·결과 글 수를 센다(파일을 안 받아도 된다).
    계획은 올렸는데 같은 해 결과 글이 없는 의회를 가려 보려는 것. 게시판에 안 올렸다고 보고서가 없다고 단정하지 말 것(재료표 조건)."""
    곳들 = json.load(open(os.path.join(ROOT, 'data', 'board_posts_council.json'), encoding='utf-8'))['곳']
    out = {}
    for laf, v in 곳들.items():
        판 = v.get('국외출장') or {}
        if 판.get('종류') != '게시판':
            continue
        해별 = {}
        for g in 판.get('글') or []:
            if not re.search(r'국외|해외|출장|연수', g['제목']):
                continue
            m = re.search(r'(20\d{2})', g['제목']) or re.search(r'(20\d{2})', g.get('날짜') or '')
            if not m:
                continue
            k = 해별.setdefault(m.group(1), {'계획': 0, '결과': 0, '기타': 0})
            k[출장갈래(g['제목'])] += 1
        if 해별:
            out[laf] = {'이름': v['이름'], '해별': 해별}
    return out


def 요약(줄모음):
    """곳·고르기마다 — 기간, 건수, 합계, 카드·현금 몫, 많이 쓴 곳 다섯(장소 이름 그대로 — 지점 이름이 다르면 따로 센다).
    판정이 아니라 회의 재료다. 경조사·격려금은 장소 칸에 사람·부서가 적힌 곳이 많다."""
    묶음 = {}
    for r in 줄모음:
        묶음.setdefault((r['곳'], r['고르기']), []).append(r)
    out = []
    for (곳, 고르기), rs in 묶음.items():
        날들 = sorted(r['날짜'] for r in rs if r['날짜'])
        장소 = {}
        for r in rs:
            if r['장소']:
                장소[r['장소']] = 장소.get(r['장소'], 0) + r['금액']
        카드 = sum(r['금액'] for r in rs if '카드' in r['방법'])
        현금 = sum(r['금액'] for r in rs if '현금' in r['방법'])
        out.append({'곳': 곳, '이름': rs[0]['이름'], '고르기': 고르기, '처음': 날들[0] if 날들 else '', '끝': 날들[-1] if 날들 else '',
                    '건수': len(rs), '합계': sum(r['금액'] for r in rs), '카드': 카드, '현금': 현금,
                    '경조사': sum(r['금액'] for r in rs if re.search(r'경조|부의|축의|조의|화환|조화', r['내용'] + r['장소'])),
                    '많이쓴곳': sorted(장소.items(), key=lambda x: -x[1])[:5]})
    return sorted(out, key=lambda x: (x['고르기'], x['곳']))


def main():
    if '--파일' in sys.argv:
        p = sys.argv[sys.argv.index('--파일') + 1]
        표들, 꼴 = 표꺼내기(open(p, 'rb').read(), os.path.splitext(p)[1][1:].lower())
        print(꼴, len(표들), '표')
        for t in 표들[:4]:
            for r in t[:8]:
                print('  ', r)
            print('   풀면:', (표풀기(t, 2026) or [])[:3])
        return
    글들 = json.load(open(os.path.join(ROOT, 'data', 'post_files.json'), encoding='utf-8'))['글']
    줄모음, 읽음, 곳별본 = [], {}, {}
    for 열쇠, 글 in 글들.items():
        if 글.get('까닭') and not 글.get('파일'):
            읽음[열쇠] = {'곳': 글['곳'], '고르기': 글['고르기'], '까닭': 글['까닭']}
            continue
        if 글['고르기'] == '국외출장':
            받은 = [f for f in 글.get('파일') or [] if f.get('sha1')]
            읽음[열쇠] = {'곳': 글['곳'], '고르기': '국외출장', '갈래': 출장갈래(글['제목']), '파일수': len(받은),
                         '쪽수': sum(쪽수(f) for f in 받은), '크기': sum(f['크기'] for f in 받은)}
            continue
        줄들, 파일들 = 글읽기(글)
        # 한 글을 열면 그 게시판 파일이 다 보이는 곳이 있다(대전 중구·유성구) — 같은 곳의 같은 건은 한 번만
        본 = 곳별본.setdefault(글['곳'], set())
        줄들 = [r for r in 줄들 if (k := (r['날짜'], r['금액'], r['장소'], r['내용'])) not in 본 and not 본.add(k)]
        for r in 줄들:
            줄모음.append({'곳': 글['곳'], '이름': 글['이름'], '고르기': 글['고르기'], '글': 글['제목'][:80], **r})
        읽음[열쇠] = {'곳': 글['곳'], '고르기': 글['고르기'], '줄': len(줄들), '합계': sum(r['금액'] for r in 줄들), '파일': 파일들}
    오늘 = datetime.date.today().isoformat()
    with gzip.open(os.path.join(ROOT, 'data', 'expense_rows.json.gz'), 'wt', encoding='utf-8') as f:
        json.dump({'만든날': 오늘, '줄': 줄모음}, f, ensure_ascii=False)
    json.dump({'만든날': 오늘, '글': 읽음}, open(os.path.join(ROOT, 'data', 'post_read.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    json.dump({'만든날': 오늘, '곳': 요약(줄모음)}, open(os.path.join(ROOT, 'data', 'expense_summary.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    출장 = 출장셈()
    json.dump({'만든날': 오늘, '곳': 출장}, open(os.path.join(ROOT, 'data', 'trip_reports.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    빈 = sum(1 for v in 출장.values() for h in v['해별'].values() if h['계획'] and not h['결과'])
    print(f'국외출장 — 의회 {len(출장)}곳 · 계획 글은 있는데 같은 해 결과 글이 없는 의회·해 {빈}개')

    # 셈 — 글 몇 건이 줄로 풀렸나, 못 푼 까닭은
    셈 = {}
    for x in 읽음.values():
        if x['고르기'] == '국외출장':
            k = f"국외출장 {x.get('갈래') or x.get('까닭', '')[:12]}" + (' · 파일 받음' if x.get('파일수') else ' · 파일 없음')
        elif x.get('줄'):
            k = '풀림'
        elif x.get('까닭'):
            k = x['까닭'][:16]
        elif any(f['까닭'].endswith('줄') for f in x.get('파일') or []):
            k = '앞 글과 같은 파일(겹쳐서 앞에서 셈)'
        else:
            k = '안 풀림: ' + ','.join(sorted({f['까닭'][:12] for f in x.get('파일') or []}) or ['파일 없음'])
        셈[k] = 셈.get(k, 0) + 1
    print(f'글 {len(읽음)}건 · 건별 줄 {len(줄모음):,}개 · 합계 {sum(r["금액"] for r in 줄모음):,}원')
    for k, v in sorted(셈.items(), key=lambda x: -x[1]):
        print(f'  {v:5}  {k}')


if __name__ == '__main__':
    main()
